from openpyxl import load_workbook
from openpyxl.styles import Font

WB = 'deliverables/Laifen-Canada-关键词与可行性分析.xlsx'
wb = load_workbook(WB)
ws = wb['01_结论与评分']
for r in range(1, ws.max_row + 1):
    a = str(ws.cell(r, 1).value or '')
    if '【最终结论】' in a:
        ws.cell(r, 2, '有条件可行 — 建议做「高低端双轨的吹风机决策站」')
        ws.cell(r, 2).font = Font(bold=True, size=11, color='006100')
    if '总体评分' in a:
        ws.cell(r, 2, '7.8 / 10  (两轮上调：5.7 → 7.3 → 7.8)')
        ws.cell(r, 2).font = Font(bold=True, size=14, color='006100')
    if '【建议打法】' in a:
        ws.cell(r, 3, '本轮更新：加入 slopehill(20%) 与 dreame(10%) 后，站点定位从「Laifen 站」升级为「高低端双轨站」')
    if '第一优先内容' in a:
        ws.cell(r, 2, '平替与对比 + 预算榜单')
        ws.cell(r, 3, '"hair dryer like dyson"/"laifen vs dyson" 收高端；"best budget hair dryer" 收 slopehill — 竞品位空缺、意图最热')
    if '站点定位' in a:
        ws.cell(r, 2, '吹风机决策站（高低端双轨）')
        ws.cell(r, 3, 'Laifen 为高端锚点 + slopehill 为走量引擎 + dreame 补充，覆盖 C$20-220 全价格带')
wb.save(WB)
print('done')
ws2 = wb['01_结论与评分']
for r in range(1, 8):
    print(r, '|', ws2.cell(r, 1).value, '||', ws2.cell(r, 2).value)
