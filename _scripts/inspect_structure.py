import glob, os, sys
import pandas as pd
import openpyxl

pd.set_option('display.width', 250)
pd.set_option('display.max_columns', 50)

files = sorted(glob.glob(os.path.join('semrush数据', '*.xlsx')))
print(f"total files: {len(files)}\n")

for f in files:
    try:
        wb = openpyxl.load_workbook(f, read_only=True)
        sheets = wb.sheetnames
        ws = wb[sheets[0]]
        rows = ws.max_row
        cols = ws.max_column
        # header row
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        print(f"=== {os.path.basename(f)}")
        print(f"    sheets={sheets} rows={rows} cols={cols}")
        print(f"    header={hdr}")
        wb.close()
    except Exception as e:
        print(f"!!! {f}: {e}")
