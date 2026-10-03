from openpyxl import load_workbook
from openpyxl.styles import Font

WB = 'deliverables/Laifen-Canada-关键词与可行性分析.xlsx'
wb = load_workbook(WB)
ws = wb['21_单笔佣金矩阵']
r = ws.max_row + 2
ws.cell(r, 1, '口径说明').font = Font(bold=True, size=10, color='1F3864')
notes = [
    '"加权均价" = 该品牌月销售额 / 月销量（销量加权，非简单平均），是估算佣金的正确基数。',
    '"单笔" = 一笔成交订单的佣金，按加权均价计算，不是单个 ASIN 的标价。',
    'Dyson 举例：加权均价 C$500.62（含 C$629.99 与 C$399.99 两款），故 @4% = C$20.02；'
    '若只算 C$399.99 那款则 @4% = C$16.00。',
    '"每+1%佣金月增" = 该品牌月销售额 × 1%，用于衡量谈判提升 1 个百分点的绝对价值。',
    '当前佣金率：Laifen 15% / slopehill 20% / dreame 10% / 其余按 Amazon 类目标准 4%。',
]
for i, t in enumerate(notes, 1):
    ws.cell(r + i, 1, f'{i}. {t}')

wb.save(WB)
print('footnote added to 21_单笔佣金矩阵')
