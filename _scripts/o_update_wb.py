"""
Append BSR-driven sheets to the deliverable workbook and update the dashboard verdict.
"""
import io, zipfile, re
import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

WB = 'deliverables/Laifen-Canada-关键词与可行性分析.xlsx'
BSR = r"C:\Users\Mloong\Downloads\BSR(Hair-Dryers(Current))-50-CA-20261002.xlsx"

# ---------- decode headers ----------
z = zipfile.ZipFile(BSR)
raw = z.read('xl/sharedStrings.xml').decode('utf-8')
S = [''.join(re.findall(r'<t[^>]*>(.*?)</t>', s, re.S)).replace('&amp;', '&').replace('&apos;', "'")
     for s in re.findall(r'<si>(.*?)</si>', raw, re.S)]
HDR = S[:71]
df = pd.read_excel(BSR, sheet_name='CA', header=0)
df.columns = HDR
for c in ['月销量', '月销售额(C$)', '价格(C$)', '评分数', '月新增评分数', '评分',
          '留评率', '大类BSR', '小类BSR', '变体数', '上架天数', 'Q&A数', '月销量增长率']:
    df[c] = pd.to_numeric(df[c], errors='coerce')

WBF = load_workbook(WB)
if '15_Amazon类目BSR' in WBF.sheetnames:
    for s in ['15_Amazon类目BSR', '16_Laifen单品与佣金', '17_品牌集中度', '18_修正收入模型']:
        if s in WBF.sheetnames:
            del WBF[s]

H_FILL = PatternFill('solid', fgColor='1F3864')
H_FONT = Font(color='FFFFFF', bold=True, size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add(name, d, widths, fmts=None, cscale=None):
    ws = WBF.create_sheet(name[:31])
    ws.append(list(d.columns))
    for c in range(1, len(d.columns) + 1):
        cell = ws.cell(1, c); cell.fill = H_FILL; cell.font = H_FONT
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BORDER
    for _, r in d.iterrows():
        ws.append(['' if (isinstance(v, float) and pd.isna(v))
                   else (v.item() if hasattr(v, 'item') else v) for v in r.tolist()])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if fmts:
        for col, f in fmts.items():
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col).number_format = f
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    if cscale:
        for col in cscale:
            L = get_column_letter(col)
            ws.conditional_formatting.add(f'{L}2:{L}{ws.max_row}',
                ColorScaleRule(start_type='min', start_color='F8696B',
                               mid_type='percentile', mid_value=50, mid_color='FFEB84',
                               end_type='max', end_color='63BE7B'))
    return ws

# ---------- 15: BSR top 50 ----------
d15 = df[['小类BSR', '品牌', 'ASIN', '商品标题', '月销量', '月销量增长率', '月销售额(C$)',
          '价格(C$)', '评分数', '月新增评分数', '评分', '留评率', '上架天数', '变体数',
          '卖家数', 'Buybox卖家', '大类BSR']].sort_values('小类BSR').copy()
d15['月销量增长率'] = (d15['月销量增长率'] * 100).round(1)
d15['留评率'] = (d15['留评率'] * 100).round(2)
d15.columns = ['小类BSR', '品牌', 'ASIN', '商品标题', '月销量', '月销量增长率%', '月销售额(C$)',
               '价格(C$)', '评分数', '月新增评分数', '评分', '留评率%', '上架天数', '变体数',
               '卖家数', 'Buybox卖家', '大类BSR']
add('15_Amazon类目BSR', d15,
    [8, 17, 13, 62, 10, 13, 14, 11, 10, 13, 8, 10, 10, 8, 8, 26, 10],
    fmts={5: '#,##0', 7: '#,##0', 8: '$#,##0.00', 9: '#,##0', 10: '#,##0',
          11: '0.0', 13: '#,##0'},
    cscale=[5, 7])

# ---------- 16: Laifen ASINs & commission ----------
la = df[df['品牌'].astype(str).str.contains('laifen', case=False, na=False)].copy()
la = la[['ASIN', '商品标题', '小类BSR', '月销量', '月销量增长率', '月销售额(C$)', '价格(C$)',
         '评分数', '月新增评分数', '评分', '上架天数', '变体数']]
la['月销量增长率'] = (la['月销量增长率'] * 100).round(1)
la['单笔佣金@15%(C$)'] = (la['价格(C$)'] * 0.15).round(2)
la['月佣金池(C$)'] = (la['月销售额(C$)'] * 0.15).round(0)
la['产品状态'] = np.where(la['月销量增长率'].fillna(0) > 5, '增长',
                  np.where(la['月销量增长率'].fillna(0) < -5, '下滑', '平稳'))
la.columns = ['ASIN', '商品标题', '小类BSR', '月销量', '月销量增长率%', '月销售额(C$)', '价格(C$)',
              '评分数', '月新增评分数', '评分', '上架天数', '变体数',
              '单笔佣金@15%(C$)', '月佣金池(C$)', '产品状态']
la = la.sort_values('月销量', ascending=False)
tot_u = la['月销量'].sum(); tot_r = la['月销售额(C$)'].sum()
aov = tot_r / tot_u
summary = pd.DataFrame([
    {'指标': 'Laifen ASIN 数', '数值': f'{len(la)} 个', '说明': '全部为 Laifen Official 自营'},
    {'指标': 'Laifen 月销量合计', '数值': f'{tot_u:,.0f} 台', '说明': '占类目 Top50 的 5.9%'},
    {'指标': 'Laifen 月销售额合计', '数值': f'C${tot_r:,.0f}', '说明': '占类目 Top50 的 13.4%'},
    {'指标': '实际客单价 AOV', '数值': f'C${aov:.2f}', '说明': '销量加权，比类目中位价 C$49.39 高 2.5 倍'},
    {'指标': '单笔 15% 佣金', '数值': f'C${aov*0.15:.2f}', '说明': '★ 这是全站最有价值的转化'},
    {'指标': '其他品牌单笔佣金 @4%', '数值': f'C${49.39*0.04:.2f}', '说明': '按类目中位价估算'},
    {'指标': '价值倍数', '数值': f'{(aov*0.15)/(49.39*0.04):.1f} 倍', '说明': '卖1台Laifen = 卖9台普通吹风机'},
    {'指标': 'Laifen 全系月佣金池', '数值': f'C${tot_r*0.15:,.0f}/月', '说明': '这是所有联盟客理论上能分到的总额'},
])
add('16_Laifen单品与佣金', la,
    [13, 60, 9, 10, 13, 14, 11, 10, 13, 8, 10, 8, 16, 14, 10],
    fmts={4: '#,##0', 6: '#,##0', 7: '$#,##0.00', 8: '#,##0', 9: '#,##0',
          10: '0.0', 11: '#,##0', 13: '$#,##0.00', 14: '#,##0'},
    cscale=[4, 6])
ws = WBF['16_Laifen单品与佣金']
r0 = ws.max_row + 3
ws.cell(r0, 1, '★ Laifen 单品经济性汇总').font = Font(bold=True, size=11, color='C00000')
for i, row in summary.iterrows():
    ws.cell(r0 + 1 + i, 1, row['指标']).font = Font(bold=True)
    ws.cell(r0 + 1 + i, 2, row['数值'])
    ws.cell(r0 + 1 + i, 4, row['说明'])

# ---------- 17: brand concentration ----------
bg = df.groupby('品牌').agg(ASIN数=('ASIN', 'size'), 月销量=('月销量', 'sum'),
                            月销售额=('月销售额(C$)', 'sum'), 均价=('价格(C$)', 'mean'),
                            评分=('评分', 'mean')).sort_values('月销量', ascending=False)
bg['销量占比%'] = (100 * bg['月销量'] / df['月销量'].sum()).round(1)
bg['销售额占比%'] = (100 * bg['月销售额'] / df['月销售额(C$)'].sum()).round(1)
bg['单笔佣金@4%'] = (bg['均价'] * 0.04).round(2)
bg = bg.reset_index().rename(columns={'品牌': '品牌', '均价': '均价(C$)', '评分': '平均评分'})
add('17_品牌集中度', bg, [20, 9, 11, 14, 11, 9, 11, 12, 12, 13],
    fmts={3: '#,##0', 4: '#,##0', 5: '$#,##0.00', 6: '0.00', 7: '0.0', 8: '0.0', 10: '$#,##0.00'},
    cscale=[3, 4])

# ---------- 18: corrected revenue model ----------
SESS = {'悲观(第1年)': 1433, '基准(第1-2年)': 4503, '乐观(第2-3年)': 8734}
rows = []
for sname, s in SESS.items():
    for cr, crl in [(0.10, '保守'), (0.15, '中性'), (0.18, '较好'), (0.25, '优秀')]:
        for cv, cvl in [(0.03, '低'), (0.05, '中'), (0.08, '较好'), (0.10, '高')]:
            pure = s * cr * cv * aov * 0.15
            blend_cpc = 0.60 * aov * 0.15 + 0.40 * 49.39 * 0.04
            blend = s * cr * cv * blend_cpc
            rows.append({'排名情景': sname, '月访问量': s,
                         '出站点击率': f'{cr:.0%} ({crl})', '成交率': f'{cv:.0%} ({cvl})',
                         '单笔佣金(C$)': round(aov * 0.15, 2),
                         '纯Laifen收入/月(C$)': round(pure),
                         '混合品牌收入/月(C$)': round(blend),
                         '纯Laifen每访问(C$)': round(pure / s, 3),
                         '混合每访问(C$)': round(blend / s, 3)})
rm = pd.DataFrame(rows)
add('18_修正收入模型', rm, [15, 11, 14, 12, 14, 20, 20, 17, 15],
    fmts={2: '#,##0', 5: '$#,##0.00', 6: '#,##0', 7: '#,##0', 8: '$#,##0.000', 9: '$#,##0.000'},
    cscale=[6, 7])

# ---------- update dashboard ----------
ws = WBF['01_结论与评分']
upd = {
    6:  ('总体评分', '7.3 / 10  (上调，原 5.7)',
         'Laifen 15% 佣金 + C$123 实测客单价，单笔佣金 C$18.47 — 是全类目最有价值的转化；$1,000/月目标仅需 3,759 访问'),
}
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == '【最终结论】':
        ws.cell(r, 2, '有条件可行 — 值得做，但要做成「Laifen 为核心的吹风机决策站」')
        ws.cell(r, 2).font = Font(bold=True, size=11, color='006100')
    if ws.cell(r, 1).value == '总体评分':
        ws.cell(r, 2, '7.3 / 10  (上调，原 5.7)')
        ws.cell(r, 2).font = Font(bold=True, size=14, color='006100')
        ws.cell(r, 3, 'Laifen 15% 佣金 + C$123 实测客单价 → 单笔佣金 C$18.47，是全类目最有价值的转化；月入 C$1,000 仅需 3,759 访问')
    if ws.cell(r, 1).value and '4. 联盟佣金结构性地薄' in str(ws.cell(r, 1).value):
        ws.cell(r, 1, '4. 佣金结构已证实【已解决】')
        ws.cell(r, 2, '利好')
        ws.cell(r, 3, 'Laifen 给 15% 佣金且客单价 C$123（类目中位价仅 C$49），单笔佣金 C$18.47 = 卖9台普通吹风机。每访问收入从假设 C$0.054 升到 C$0.266（4.9倍）')
        ws.cell(r, 2).font = Font(bold=True, color='006100')
    if ws.cell(r, 1).value and '继续投入的条件' in str(ws.cell(r, 1).value):
        ws.cell(r, 3, '已部分满足：15% 佣金已确认。剩下需要验证的是「出站点击率」与「成交率」两个真实转化指标')
    if ws.cell(r, 1).value and '最大不确定性' in str(ws.cell(r, 1).value):
        ws.cell(r, 3, '出站点击率与 Amazon 成交率目前仍是假设（用 18% x 8% 建模）。需要上线后用 Amazon Associates 后台实测校准')
    if ws.cell(r, 1).value == '【收入模型】':
        ws.cell(r, 2, 'base 情景（已用实测客单价与佣金率修正）')

r = ws.max_row + 2
ws.cell(r, 1, '【Amazon.ca 实测数据补充】').font = Font(bold=True, size=11, color='1F3864')
for c in range(1, 4):
    ws.cell(r, c).fill = PatternFill('solid', fgColor='D9E2F3')
extra = [
    ('类目 Top50 月销量', '33,507 台', 'Amazon.ca 吹风机类目'),
    ('类目 Top50 月销售额', 'C$1,809,118', '平均客单价 C$53.99'),
    ('Laifen 月销量', f'{tot_u:,.0f} 台 (5.9%)', '6 个 ASIN，全部 Laifen Official 自营'),
    ('Laifen 月销售额', f'C${tot_r:,.0f} (13.4%)', '用 5.9% 销量拿 13.4% 销售额'),
    ('Laifen 单品冠军', 'B0GWHF4SHD Laifen Air', '534台/月，C$99.99，上架仅126天'),
    ('Laifen 增长最快', 'B0GWHGTBDY Air Diffuser', '+85条/月新增评分，上架122天'),
    ('Laifen 衰退款', 'B0D141Q8ZF Swift gen1', '-39% 且评分仅4.1 — 不建议主推'),
    ('加权单笔佣金(混合品牌)', f'C${0.60*aov*0.15+0.40*49.39*0.04:.2f}',
     '假设60%出站去Laifen、40%去其他品牌'),
]
for i, (a, b, c) in enumerate(extra, start=1):
    ws.cell(r + i, 1, a); ws.cell(r + i, 2, b); ws.cell(r + i, 3, c)

WBF.save(WB)
print('updated:', WB)
print('sheets:', WBF.sheetnames)
print(f"\nAOV = C${aov:.2f}   单笔佣金@15% = C${aov*0.15:.2f}")
