"""Add commission-pool + product-mix sheets and update the verdict."""
import io
import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

WB = 'deliverables/Laifen-Canada-关键词与可行性分析.xlsx'
L, SL, DR = 18.47, 9.84, 16.00          # 单笔佣金
W = 0.35 * L + 0.55 * SL + 0.10 * DR    # 均衡组合加权

WBF = load_workbook(WB)
for s in ['19_品牌佣金池', '20_产品组合策略']:
    if s in WBF.sheetnames:
        del WBF[s]

H_FILL = PatternFill('solid', fgColor='1F3864')
H_FONT = Font(color='FFFFFF', bold=True, size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add(name, d, widths, fmts=None, cscale=None):
    ws = WBF.create_sheet(name[:31])
    ws.append(list(d.columns))
    for c in range(1, len(d.columns) + 1):
        cell = ws.cell(1, c); cell.fill = H_FILL; cell.font = H_FONT
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BORDER
    for _, r in d.iterrows():
        ws.append(['' if (isinstance(v, float) and pd.isna(v)) else v for v in r.tolist()])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if fmts:
        for col, f in fmts.items():
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col).number_format = f
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    if cscale:
        for col in cscale:
            Lc = get_column_letter(col)
            ws.conditional_formatting.add(f'{Lc}2:{Lc}{ws.max_row}',
                ColorScaleRule(start_type='min', start_color='F8696B',
                               mid_type='percentile', mid_value=50, mid_color='FFEB84',
                               end_type='max', end_color='63BE7B'))
    return ws


# ---------------- 19: commission pool by brand ----------------
pool = pd.read_pickle('analysis/brand_commission.pkl')
pool = pool.rename(columns={'品牌': '品牌', 'ASIN数': 'ASIN数', '月销量': '月销量',
                            '月销售额': '月销售额(C$)', '加权均价': '加权均价(C$)',
                            '佣金率': '佣金率', '单笔佣金': '单笔佣金(C$)',
                            '月佣金池': '月佣金池(C$)'})
pool['是否自定义佣金'] = np.where(pool['品牌'].isin(['Laifen', 'slopehill', 'dreame']), '★是', '')
pool['占总佣金池%'] = (100 * pool['月佣金池(C$)'] / pool['月佣金池(C$)'].sum()).round(1)
pool = pool[['品牌', '是否自定义佣金', 'ASIN数', '月销量', '月销售额(C$)', '加权均价(C$)',
             '佣金率', '单笔佣金(C$)', '月佣金池(C$)', '占总佣金池%']]
ws19 = add('19_品牌佣金池', pool,
           [17, 14, 8, 10, 14, 12, 9, 13, 13, 12],
           fmts={4: '#,##0', 5: '#,##0', 6: '$#,##0.00', 7: '0%', 8: '$#,##0.00', 9: '#,##0'},
           cscale=[9])

r0 = ws19.max_row + 3
ws19.cell(r0, 1, '★ 三个自定义佣金品牌 —— 策略含义').font = Font(bold=True, size=11, color='C00000')
notes = [
    ('slopehill 月佣金池', 'C$98,956 (55.0%)', '类目最大池子，但 Buy Box 由第三方 Stellartech 掌控，3个ASIN共2-5个跟卖'),
    ('Laifen 月佣金池', 'C$36,393 (20.2%)', '单笔最值钱 C$18.47；Buy Box 100% 自控（Laifen Official），结构最稳'),
    ('dreame 月佣金池', 'C$2,608 (1.5%)', '量太小，定位为补充高客单 SKU，不做主推'),
    ('三品牌合计', 'C$137,957 = 类目 Top50 佣金池的 76.7%', '只做 Laifen 会放弃 80% 可赚佣金'),
    ('关键结构差异', 'Laifen 自控 / slopehill 第三方', 'Laifen 链接稳定性远优于 slopehill，佣金不会因 Buy Box 易主而流失'),
]
for i, (a, b, c) in enumerate(notes, 1):
    ws19.cell(r0 + i, 1, a).font = Font(bold=True)
    ws19.cell(r0 + i, 2, b)
    ws19.cell(r0 + i, 4, c)

# ---------------- 20: product mix strategy ----------------
rows = []
for name, a, b, c in [
    ('只推 Laifen', 1.00, 0.00, 0.00),
    ('Laifen 高端为主', 0.60, 0.35, 0.05),
    ('★ 高端+走量均衡（推荐）', 0.35, 0.55, 0.10),
    ('走量为主 slopehill', 0.15, 0.75, 0.10),
    ('只推 slopehill', 0.00, 1.00, 0.00),
]:
    w = a * L + b * SL + c * DR
    need1k = 1000 / (0.18 * 0.08 * w)
    rows.append({'内容策略': name, 'Laifen占比': a, 'slopehill占比': b, 'dreame占比': c,
                 '加权单笔佣金(C$)': round(w, 2), '相对纯Laifen': w / L,
                 '月入C$1,000所需访问': round(need1k),
                 '基准情景4503访问月收入(C$)': round(4503 * 0.18 * 0.08 * w)})
mix = pd.DataFrame(rows)
ws20 = add('20_产品组合策略', mix, [26, 11, 13, 11, 16, 13, 18, 22],
           fmts={2: '0%', 3: '0%', 4: '0%', 5: '$#,##0.00', 6: '0%', 7: '#,##0', 8: '#,##0'},
           cscale=[8])

r0 = ws20.max_row + 3
ws20.cell(r0, 1, '★ 为什么推荐「均衡」组合').font = Font(bold=True, size=11, color='C00000')
for i, txt in enumerate([
    '均衡组合单笔佣金 C$13.48，比纯 Laifen (C$18.47) 低 27%',
    '但可覆盖 C$20-220 全价格带搜索意图，可寻址流量远大于只做高端',
    'slopehill 本身就是最大的「Dyson 平替」（C$42.99 vs Dyson C$400+），与站点「平替」定位完全吻合',
    '只做 Laifen 需 3,760 访问达 C$1,000；均衡组合需 5,153 访问 —— 多 1,393 访问换更宽、更抗风险的结构',
    '操作要点：slopehill 20% 必须是卖家(Stellartech等)直连协议；若走 Amazon Associates 只能拿类目标准 4%',
], 1):
    ws20.cell(r0 + i, 1, f'{i}. {txt}')

# ---------------- update dashboard ----------------
ws = WBF['01_结论与评分']
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v == '总体评分':
        ws.cell(r, 2, '7.8 / 10  (上调，原 7.3)')
        ws.cell(r, 2).font = Font(bold=True, size=14, color='006100')
        ws.cell(r, 3, '三个品牌自定义佣金（Laifen 15%/slopehill 20%/dreame 10%）覆盖类目 76.7% 佣金池；'
                      '均衡组合加权单笔佣金 C$13.48，月入 C$1,000 需 5,153 访问')
    if v and '站点定位' in str(v):
        ws.cell(r, 2, '吹风机决策站（高低端双轨）')
        ws.cell(r, 3, '以 Laifen 为高端锚点 + slopehill 为走量引擎，覆盖 C$20-220 全价格带')
r = ws.max_row + 2
ws.cell(r, 1, '【三个自定义佣金品牌 —— 佣金池实测】').font = Font(bold=True, size=11, color='1F3864')
for c in range(1, 4):
    ws.cell(r, c).fill = PatternFill('solid', fgColor='D9E2F3')
extra = [
    ('slopehill 20%', '单笔 C$9.84 / 池 C$98,956', '类目最大池子(55%)，C$49均价，月销 10,052 台'),
    ('Laifen 15%', '单笔 C$18.47 / 池 C$36,393', '单笔最值钱，C$123均价，Buy Box 100% 自控'),
    ('dreame 10%', '单笔 C$16.00 / 池 C$2,608', 'C$160均价但量小，作补充 SKU'),
    ('三品牌合计', 'C$137,957/月 = 76.7%', '占类目 Top50 总佣金池的 76.7%'),
    ('推荐加权佣金', 'C$13.48 / 笔', '35% Laifen + 55% slopehill + 10% dreame'),
]
for i, (a, b, c) in enumerate(extra, 1):
    ws.cell(r + i, 1, a); ws.cell(r + i, 2, b); ws.cell(r + i, 3, c)

WBF.save(WB)
print('updated. sheets:', WBF.sheetnames)
print(f'weighted commission = C${W:.2f}/sale')
